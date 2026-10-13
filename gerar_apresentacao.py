from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT = "apresentacao_titanic.pptx"


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.7))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.size = Pt(28)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(20, 45, 78)
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.65), Inches(1.05), Inches(11.8), Inches(0.45))
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = subtitle
        subtitle_frame.paragraphs[0].font.size = Pt(15)
        subtitle_frame.paragraphs[0].font.color.rgb = RGBColor(80, 90, 100)


def add_bullets(slide, items, top=1.5, font_size=22):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(top), Inches(11.4), Inches(5.3))
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor(35, 45, 55)
        paragraph.space_after = Pt(14)


presentation = Presentation()
presentation.slide_width = Inches(13.333)
presentation.slide_height = Inches(7.5)
blank_layout = presentation.slide_layouts[6]

slide = presentation.slides.add_slide(blank_layout)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(232, 245, 243)
box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.8), Inches(1.4))
frame = box.text_frame
frame.text = "Métricas de classificação: Titanic"
frame.paragraphs[0].font.size = Pt(38)
frame.paragraphs[0].font.bold = True
frame.paragraphs[0].font.color.rgb = RGBColor(20, 90, 90)
frame.paragraphs[0].alignment = PP_ALIGN.CENTER
subtitle = slide.shapes.add_textbox(Inches(1.3), Inches(3.6), Inches(10.8), Inches(0.8))
subtitle.text_frame.text = "Previsão de sobrevivência com regressão logística"
subtitle.text_frame.paragraphs[0].font.size = Pt(24)
subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(55, 70, 75)
subtitle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

slide = presentation.slides.add_slide(blank_layout)
add_title(slide, "1. Problema e objetivo")
add_bullets(slide, [
    "Objetivo: prever se um passageiro sobreviveu ao naufrágio.",
    "Variável-alvo: survived (0 = não sobreviveu; 1 = sobreviveu).",
    "Avaliar o modelo com matriz de confusão, acurácia, precisão e recall.",
])

slide = presentation.slides.add_slide(blank_layout)
add_title(slide, "2. Base de dados")
add_bullets(slide, [
    "Base Titanic pública, com 891 registros e 15 colunas.",
    "Variáveis utilizadas: classe, sexo, idade, tarifa, familiares e porto de embarque.",
    "A classe 0 representa 61,62% dos registros e a classe 1 representa 38,38%.",
    "Fonte: Seaborn Titanic dataset.",
])

slide = presentation.slides.add_slide(blank_layout)
add_title(slide, "3. Tratamento dos dados")
add_bullets(slide, [
    "Idade: valores ausentes preenchidos pela mediana.",
    "Embarque: valores ausentes preenchidos pela categoria mais frequente.",
    "Variáveis categóricas transformadas com one-hot encoding.",
    "Variáveis numéricas padronizadas e dados divididos em 80% treino e 20% teste.",
])

slide = presentation.slides.add_slide(blank_layout)
add_title(slide, "4. Modelo utilizado")
add_bullets(slide, [
    "Modelo: regressão logística.",
    "Escolha: modelo simples, rápido e adequado para classificação binária.",
    "Treino: 712 registros | Teste: 179 registros.",
    "A semente aleatória foi fixada em 42 para permitir reprodução dos resultados.",
])

slide = presentation.slides.add_slide(blank_layout)
add_title(slide, "5. Resultados")
add_bullets(slide, [
    "Acurácia: 80,45%",
    "Precisão: 79,31%",
    "Recall: 66,67%",
    "F1-score: 72,44%",
], top=1.45, font_size=25)

slide = presentation.slides.add_slide(blank_layout)
add_title(slide, "6. Matriz de confusão")
add_bullets(slide, [
    "Verdadeiros negativos: 98",
    "Falsos positivos: 12",
    "Falsos negativos: 23",
    "Verdadeiros positivos: 46",
    "O modelo acertou mais casos de não sobrevivência do que casos de sobrevivência.",
], top=1.45, font_size=22)

slide = presentation.slides.add_slide(blank_layout)
add_title(slide, "7. Avaliação final")
add_bullets(slide, [
    "O desempenho geral foi bom, com 80,45% de acurácia.",
    "A precisão de 79,31% indica que a maioria das previsões de sobrevivência estava correta.",
    "O recall de 66,67% mostra que 23 sobreviventes não foram identificados.",
    "O modelo é útil como análise inicial, mas pode melhorar com mais variáveis e comparação com outros algoritmos.",
])

slide = presentation.slides.add_slide(blank_layout)
add_title(slide, "Conclusão")
add_bullets(slide, [
    "O pipeline transformou os dados e treinou um classificador reproduzível.",
    "As métricas mostram um modelo equilibrado, porém com espaço para melhorar a identificação dos sobreviventes.",
    "Notebook, base CSV e apresentação compõem a entrega do projeto.",
])

presentation.save(OUTPUT)
print(f"Arquivo criado: {OUTPUT}")
